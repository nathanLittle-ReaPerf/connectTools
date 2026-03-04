#!/usr/bin/env python3
"""cid_journey.py — Render a Cytoscape.js caller journey map from a CID_Search xlsx."""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print(
        "openpyxl is required.  Install with:  pip install openpyxl --user",
        file=sys.stderr,
    )
    sys.exit(1)

# ── Constants ─────────────────────────────────────────────────────────────────

EXPECTED_COLUMNS = [
    "@timestamp", "ContactId", "ContactFlowName", "ContactFlowModuleType",
    "Attribute", "Value", "Check", "Results", "Prompt", "Operation",
    "Function", "Parameters", "External_Results",
]

MODULE_LABELS: dict[str, str] = {
    "PlayPrompt":              "Play Prompt",
    "MessageParticipant":      "Play Message",
    "CheckAttribute":          "Check Attribute",
    "GetUserInput":            "Get Input",
    "StoreCustomerInput":      "Store Input",
    "InvokeExternalResource":  "Lambda",
    "InvokeLambdaFunction":    "Lambda",
    "SetAttributes":           "Set Attribute",
    "UpdateContactAttributes": "Update Attrs",
    "SetLoggingBehavior":      "Set Logging",
    "SetContactFlow":          "Set Flow",
    "Transfer":                "Transfer",
    "TransferToQueue":         "Transfer to Queue",
    "TransferContactToQueue":  "Transfer to Queue",
    "TransferToAgent":         "Transfer to Agent",
    "Disconnect":              "Disconnect",
    "DisconnectParticipant":   "Disconnect",
    "CheckHoursOfOperation":   "Check Hours",
    "CheckAgentStatus":        "Check Agent",
    "WaitForCustomerInput":    "Wait for Input",
    "StartMediaStreaming":      "Start Streaming",
    "StopMediaStreaming":       "Stop Streaming",
}

DECISION_TYPES = {
    "CheckAttribute", "GetUserInput", "StoreCustomerInput",
    "CheckHoursOfOperation", "CheckAgentStatus", "WaitForCustomerInput",
}

TERMINAL_TYPES = {
    "Disconnect", "Transfer", "TransferToQueue", "TransferContactToQueue",
    "TransferToAgent", "DisconnectParticipant",
}

INVOKE_TYPES = {"InvokeExternalResource", "InvokeLambdaFunction"}

# ── Helpers ───────────────────────────────────────────────────────────────────

def _parse_ts(s: str) -> dt.datetime | None:
    """Parse a CloudWatch @timestamp into a datetime, or None on failure."""
    if not s:
        return None
    s = s.strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S",
    ):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _delta_s(t1: dt.datetime | None, t2: dt.datetime | None) -> int | None:
    if t1 is None or t2 is None:
        return None
    return max(0, int((t2 - t1).total_seconds()))


def _classify(mod_type: str) -> str:
    if mod_type in DECISION_TYPES:
        return "decision"
    if mod_type in TERMINAL_TYPES:
        return "terminal"
    if mod_type in INVOKE_TYPES:
        return "invoke"
    return "default"


def _trunc(s: str, n: int) -> str:
    s = (s or "").strip()
    return s[:n] + "\u2026" if len(s) > n else s


def _pick_detail(row: dict, mod_type: str) -> str:
    if mod_type in ("PlayPrompt", "MessageParticipant", "GetUserInput"):
        return _trunc(row.get("Prompt", ""), 40)
    if mod_type in ("SetAttributes", "UpdateContactAttributes"):
        attr = row.get("Attribute", "")
        val  = row.get("Value", "")
        return _trunc(f"{attr}={val}", 36) if attr else ""
    if mod_type == "CheckAttribute":
        attr  = row.get("Attribute", "")
        check = row.get("Check", "")
        res   = row.get("Results", "")
        parts = [attr]
        if check:
            parts.append(f"={check}")
        if res:
            parts.append(f" \u2192{res}")
        return _trunc("".join(parts), 36)
    if mod_type in INVOKE_TYPES:
        return _trunc(row.get("Function", "") or row.get("Operation", ""), 36)
    return _trunc(row.get("Results", ""), 36)


def _make_label(row: dict) -> str:
    mod_type  = row.get("ContactFlowModuleType", "")
    flow_name = row.get("ContactFlowName", "")
    line1  = MODULE_LABELS.get(mod_type, mod_type) or "?"
    line2  = _trunc(flow_name, 28)
    detail = _pick_detail(row, mod_type)
    parts  = [line1]
    if line2:
        parts.append(line2)
    if detail:
        parts.append(detail)
    return "\n".join(parts)


# ── Data loading ──────────────────────────────────────────────────────────────

def load_xlsx(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        wb.close()
        return []

    headers = [str(h).strip() if h is not None else "" for h in raw_headers]
    missing = [c for c in EXPECTED_COLUMNS if c not in headers]
    if missing:
        print(f"  Warning: expected columns not found: {', '.join(missing)}", file=sys.stderr)

    result = []
    for raw_row in rows_iter:
        if all(v is None for v in raw_row):
            continue
        row: dict[str, str] = {}
        for i, h in enumerate(headers):
            if h:
                v = raw_row[i] if i < len(raw_row) else None
                row[h] = "" if v is None else str(v).strip()
        result.append(row)

    wb.close()
    return result


# ── Graph building ─────────────────────────────────────────────────────────────

def build_elements(rows: list[dict]) -> list[dict]:
    """Return a flat Cytoscape elements list (nodes then edges)."""
    node_data_list = []
    visit_counts: dict[tuple, int] = {}
    timestamps: list[dt.datetime | None] = []

    for i, row in enumerate(rows):
        mod_type  = row.get("ContactFlowModuleType", "")
        flow_name = row.get("ContactFlowName", "")
        ts        = _parse_ts(row.get("@timestamp", ""))
        timestamps.append(ts)

        pair = (flow_name, mod_type)
        visit_counts[pair] = visit_counts.get(pair, 0) + 1
        repeated = visit_counts[pair] > 1

        ntype = "start" if i == 0 else _classify(mod_type)
        node_data_list.append({
            "id":       f"n{i}",
            "label":    _make_label(row),
            "ntype":    ntype,
            "repeated": repeated,
            "ts":       ts.isoformat() if ts else "",
            "row":      dict(row),
        })

    # Inject timing delta into each node's row dict (shown in detail panel)
    for i in range(1, len(node_data_list)):
        d = _delta_s(timestamps[i - 1], timestamps[i])
        if d is not None:
            node_data_list[i]["row"]["_delta_from_prev"] = f"{d}s"

    # Build flat elements list
    elements = [{"data": nd} for nd in node_data_list]

    for i in range(1, len(node_data_list)):
        prev_mod   = rows[i - 1].get("ContactFlowModuleType", "")
        edge_label = rows[i - 1].get("Results", "") if prev_mod in DECISION_TYPES else ""
        elements.append({"data": {
            "id":     f"e{i}",
            "source": f"n{i - 1}",
            "target": f"n{i}",
            "label":  edge_label,
        }})

    return elements


# ── HTML template ─────────────────────────────────────────────────────────────
# Sentinel placeholders avoid escaping thousands of JS braces.

_HTML = r"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>__TITLE__</title>
  <style>
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
      background: #f0f2f5; height: 100vh; display: flex; flex-direction: column;
    }

    /* Header */
    header {
      padding: 8px 16px; background: white; border-bottom: 1px solid #ddd;
      display: flex; align-items: center; gap: 12px; flex-shrink: 0;
      position: relative; z-index: 10;
    }
    h1 { font-size: 1rem; color: #333; flex: 1; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
    .controls { display: flex; gap: 6px; align-items: center; flex-shrink: 0; }
    button { padding: 4px 12px; border: 1px solid #ccc; border-radius: 4px; background: white; cursor: pointer; font-size: 0.9rem; }
    button:hover { background: #f0f0f0; }
    #btn-colors.active { background: #E3F2FD; border-color: #90CAF9; }
    #zoom-label { font-size: 0.85rem; color: #666; min-width: 46px; text-align: center; }

    /* Canvas */
    #cy { flex: 1; }

    /* Shared panel */
    .panel-header {
      padding: 10px 14px; font-weight: 600; font-size: 0.9rem;
      border-bottom: 1px solid #eee; display: flex;
      align-items: center; justify-content: space-between; flex-shrink: 0;
    }
    .panel-header button { padding: 2px 8px; font-size: 1rem; border: none; background: none; color: #888; cursor: pointer; }
    .panel-header button:hover { background: #f5f5f5; color: #333; }
    .panel-body { overflow-y: auto; flex: 1; padding: 12px 14px; }

    /* Detail panel (left) */
    #detail-panel {
      position: fixed; left: -330px; top: 0; width: 318px; bottom: 0;
      background: white; border-right: 1px solid #ddd;
      box-shadow: 4px 0 16px rgba(0,0,0,0.1);
      transition: left 0.22s ease; z-index: 300;
      display: flex; flex-direction: column;
    }
    #detail-panel.open { left: 0; }
    .hint { color: #aaa; font-size: 0.85rem; padding: 4px 0; }
    .detail-table { width: 100%; border-collapse: collapse; font-size: 0.8rem; }
    .detail-table tr:hover td { background: #fafafa; }
    .detail-table td { padding: 4px 6px; border-bottom: 1px solid #f0f0f0; vertical-align: top; }
    .col-key { font-weight: 600; color: #555; width: 42%; white-space: nowrap; }
    .col-val { word-break: break-all; color: #333; }

    /* Color panel (right) */
    #color-panel {
      position: fixed; right: -260px; top: 0; width: 248px; bottom: 0;
      background: white; border-left: 1px solid #ddd;
      box-shadow: -4px 0 16px rgba(0,0,0,0.1);
      transition: right 0.22s ease; z-index: 300;
      display: flex; flex-direction: column;
    }
    #color-panel.open { right: 0; }
    .section-title {
      font-size: 0.75rem; font-weight: 600; text-transform: uppercase;
      color: #888; letter-spacing: 0.05em; margin: 12px 0 8px;
    }
    .section-title:first-child { margin-top: 0; }
    .preset-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 6px; margin-bottom: 4px; }
    .preset-btn {
      padding: 6px 4px; font-size: 0.8rem; border-radius: 4px;
      border: 1px solid #ddd; background: white; cursor: pointer;
      display: flex; align-items: center; justify-content: center; gap: 6px;
    }
    .preset-btn:hover { background: #f5f5f5; }
    .preset-btn .dots { display: flex; gap: 2px; }
    .preset-btn .dot { width: 8px; height: 8px; border-radius: 50%; flex-shrink: 0; }
    .color-row {
      display: flex; align-items: center; gap: 8px;
      padding: 5px 0; border-bottom: 1px solid #f5f5f5;
    }
    .color-row:last-child { border-bottom: none; }
    .swatch { width: 20px; height: 20px; border-radius: 4px; border: 1px solid rgba(0,0,0,0.15); flex-shrink: 0; }
    .color-row label { flex: 1; font-size: 0.85rem; color: #444; }
    .color-row input[type="color"] { width: 32px; height: 24px; padding: 1px; border: 1px solid #ccc; border-radius: 4px; cursor: pointer; background: white; }
  </style>
</head>
<body>
  <header>
    <h1>__TITLE__</h1>
    <div class="controls">
      <button id="btn-detail">Details</button>
      <button id="btn-zoom-in">+</button>
      <span id="zoom-label">100%</span>
      <button id="btn-zoom-out">−</button>
      <button id="btn-reset">Fit</button>
      <button id="btn-colors">Colors</button>
    </div>
  </header>

  <div id="cy"></div>

  <!-- Detail panel (left) -->
  <div id="detail-panel">
    <div class="panel-header">
      <span id="detail-title">Node Details</span>
      <button id="btn-close-detail" title="Close">×</button>
    </div>
    <div class="panel-body">
      <div id="detail-body"><p class="hint">Click a node to see its log details.</p></div>
    </div>
  </div>

  <!-- Color panel (right) -->
  <div id="color-panel">
    <div class="panel-header">
      <span>Colors</span>
      <button id="btn-close-panel" title="Close">×</button>
    </div>
    <div class="panel-body">
      <div class="section-title">Presets</div>
      <div class="preset-grid">
        <button class="preset-btn" data-theme="default">
          <span class="dots">
            <span class="dot" style="background:#2E7D32"></span>
            <span class="dot" style="background:#E3F2FD"></span>
            <span class="dot" style="background:#1565C0"></span>
            <span class="dot" style="background:#B71C1C"></span>
          </span>Default
        </button>
        <button class="preset-btn" data-theme="dark">
          <span class="dots">
            <span class="dot" style="background:#1B5E20"></span>
            <span class="dot" style="background:#1A237E"></span>
            <span class="dot" style="background:#4527A0"></span>
            <span class="dot" style="background:#7F0000"></span>
          </span>Dark
        </button>
        <button class="preset-btn" data-theme="pastel">
          <span class="dots">
            <span class="dot" style="background:#81C784"></span>
            <span class="dot" style="background:#BBDEFB"></span>
            <span class="dot" style="background:#CE93D8"></span>
            <span class="dot" style="background:#EF9A9A"></span>
          </span>Pastel
        </button>
        <button class="preset-btn" data-theme="mono">
          <span class="dots">
            <span class="dot" style="background:#424242"></span>
            <span class="dot" style="background:#F5F5F5"></span>
            <span class="dot" style="background:#757575"></span>
            <span class="dot" style="background:#212121"></span>
          </span>Mono
        </button>
      </div>

      <div class="section-title">Node Colors</div>
      <div class="color-row">
        <div class="swatch" id="sw-start"></div>
        <label>Start</label>
        <input type="color" id="c-start">
      </div>
      <div class="color-row">
        <div class="swatch" id="sw-action"></div>
        <label>Action</label>
        <input type="color" id="c-action">
      </div>
      <div class="color-row">
        <div class="swatch" id="sw-decision"></div>
        <label>Decision</label>
        <input type="color" id="c-decision">
      </div>
      <div class="color-row">
        <div class="swatch" id="sw-invoke"></div>
        <label>Lambda</label>
        <input type="color" id="c-invoke">
      </div>
      <div class="color-row">
        <div class="swatch" id="sw-terminal"></div>
        <label>Terminal</label>
        <input type="color" id="c-terminal">
      </div>

      <div class="section-title">Edge</div>
      <div class="color-row">
        <div class="swatch" id="sw-edge"></div>
        <label>Edges</label>
        <input type="color" id="c-edge">
      </div>
    </div>
  </div>

  <script src="https://cdn.jsdelivr.net/npm/cytoscape@3/dist/cytoscape.min.js"></script>
  <script src="https://cdn.jsdelivr.net/npm/dagre@0.8.5/dist/dagre.min.js"></script>
  <script src="https://cdn.jsdelivr.net/npm/cytoscape-dagre@2.5.0/cytoscape-dagre.js"></script>
  <script>
    const COLUMNS = __COL_HEADERS_JSON__;

    const cy = cytoscape({
      container: document.getElementById('cy'),
      elements:  __ELEMENTS_JSON__,
      wheelSensitivity: 0.3,
      style: [
        {
          selector: 'node',
          style: {
            'label': 'data(label)', 'text-valign': 'center', 'text-halign': 'center',
            'text-wrap': 'wrap', 'text-max-width': '160px',
            'font-size': '11px', 'font-family': 'Arial, sans-serif',
            'width': 'label', 'height': 'label', 'padding': '10px',
            'shape': 'roundrectangle', 'border-width': 1,
            'background-color': '#E3F2FD', 'border-color': '#90CAF9', 'color': '#333',
          }
        },
        {
          selector: 'node[ntype="start"]',
          style: {
            'shape': 'ellipse', 'font-weight': 'bold',
            'background-color': '#2E7D32', 'border-color': '#1B5E20', 'color': 'white',
          }
        },
        {
          selector: 'node[ntype="decision"]',
          style: {
            'shape': 'diamond', 'text-max-width': '100px', 'padding': '20px',
            'background-color': '#1565C0', 'border-color': '#0D47A1', 'color': 'white',
          }
        },
        {
          selector: 'node[ntype="invoke"]',
          style: {
            'background-color': '#E65100', 'border-color': '#BF360C', 'color': 'white',
          }
        },
        {
          selector: 'node[ntype="terminal"]',
          style: {
            'shape': 'ellipse',
            'background-color': '#B71C1C', 'border-color': '#7F0000', 'color': 'white',
          }
        },
        {
          selector: 'node[?repeated]',
          style: { 'border-style': 'dashed', 'border-width': 2 }
        },
        {
          selector: 'node:selected',
          style: { 'border-color': '#FF6F00', 'border-width': 3, 'border-style': 'solid' }
        },
        {
          selector: 'edge',
          style: {
            'label': 'data(label)', 'font-size': '9px', 'font-family': 'Arial, sans-serif',
            'color': '#444', 'text-background-color': 'white',
            'text-background-opacity': 0.9, 'text-background-padding': '2px',
            'curve-style': 'bezier', 'target-arrow-shape': 'triangle',
            'arrow-scale': 0.8, 'line-color': '#aaa', 'target-arrow-color': '#aaa',
            'width': 1.5,
          }
        }
      ],
      layout: {
        name: 'dagre', rankDir: 'TB',
        nodeSep: 60, rankSep: 80, edgeSep: 20, padding: 40,
        animate: false, fit: true,
      }
    });

    // ── Zoom controls ──────────────────────────────────────────────────────────
    const fitZoom   = cy.zoom();
    const zoomLabel = document.getElementById('zoom-label');
    function updateLabel() { zoomLabel.textContent = Math.round(cy.zoom() / fitZoom * 100) + '%'; }
    cy.on('zoom', updateLabel);
    updateLabel();
    function fit() { cy.fit(undefined, 40); updateLabel(); }
    function zoomBy(f) { cy.zoom({ level: cy.zoom() * f, renderedPosition: { x: cy.width()/2, y: cy.height()/2 } }); }
    document.getElementById('btn-zoom-in') .addEventListener('click', () => zoomBy(1.25));
    document.getElementById('btn-zoom-out').addEventListener('click', () => zoomBy(0.8));
    document.getElementById('btn-reset')   .addEventListener('click', fit);
    window.addEventListener('resize', fit);

    // ── Detail panel ───────────────────────────────────────────────────────────
    const detailPanel = document.getElementById('detail-panel');

    function esc(s) {
      return String(s ?? '').replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
    }

    function showDetail(node) {
      const row = node.data('row') || {};
      const shown = new Set();
      let html = '<table class="detail-table"><tbody>';
      for (const col of COLUMNS) {
        const v = row[col];
        if (v !== undefined && v !== '') {
          html += `<tr><td class="col-key">${esc(col)}</td><td class="col-val">${esc(v)}</td></tr>`;
          shown.add(col);
        }
      }
      // Extra synthetic keys (e.g. _delta_from_prev)
      for (const [k, v] of Object.entries(row)) {
        if (!shown.has(k) && v !== '') {
          html += `<tr><td class="col-key">${esc(k)}</td><td class="col-val">${esc(v)}</td></tr>`;
        }
      }
      html += '</tbody></table>';
      document.getElementById('detail-title').textContent = node.data('label').split('\n')[0];
      document.getElementById('detail-body').innerHTML = html;
      detailPanel.classList.add('open');
    }

    cy.on('tap', 'node', evt => showDetail(evt.target));
    cy.on('tap', evt => { if (evt.target === cy) detailPanel.classList.remove('open'); });

    document.getElementById('btn-detail').addEventListener('click', () => {
      const sel = cy.$('node:selected');
      if (sel.length) showDetail(sel[0]);
      else detailPanel.classList.toggle('open');
    });
    document.getElementById('btn-close-detail').addEventListener('click', () => {
      detailPanel.classList.remove('open');
    });

    // ── Color panel ────────────────────────────────────────────────────────────
    const THEMES = {
      default: { start:'#2E7D32', action:'#E3F2FD', decision:'#1565C0', invoke:'#E65100', terminal:'#B71C1C', edge:'#aaaaaa' },
      dark:    { start:'#1B5E20', action:'#1A237E', decision:'#4527A0', invoke:'#BF360C', terminal:'#7F0000', edge:'#555555' },
      pastel:  { start:'#81C784', action:'#BBDEFB', decision:'#CE93D8', invoke:'#FFCC80', terminal:'#EF9A9A', edge:'#BDBDBD' },
      mono:    { start:'#424242', action:'#F5F5F5', decision:'#757575', invoke:'#9E9E9E', terminal:'#212121', edge:'#9E9E9E' },
    };

    function textColor(hex) {
      const [r,g,b] = [1,3,5].map(i => parseInt(hex.slice(i,i+2),16));
      return (0.299*r + 0.587*g + 0.114*b) / 255 > 0.5 ? '#333333' : '#ffffff';
    }
    function darken(hex, pct=0.25) {
      return '#' + [1,3,5].map(i =>
        Math.round(Math.max(0, parseInt(hex.slice(i,i+2),16) * (1-pct))).toString(16).padStart(2,'0')
      ).join('');
    }

    function applyColors(c) {
      cy.batch(() => {
        cy.nodes('[ntype="start"]')   .style({ 'background-color':c.start,    'border-color':darken(c.start),    'color':textColor(c.start)    });
        cy.nodes('[ntype="default"]') .style({ 'background-color':c.action,   'border-color':darken(c.action),   'color':textColor(c.action)   });
        cy.nodes('[ntype="decision"]').style({ 'background-color':c.decision, 'border-color':darken(c.decision), 'color':textColor(c.decision) });
        cy.nodes('[ntype="invoke"]')  .style({ 'background-color':c.invoke,   'border-color':darken(c.invoke),   'color':textColor(c.invoke)   });
        cy.nodes('[ntype="terminal"]').style({ 'background-color':c.terminal, 'border-color':darken(c.terminal), 'color':textColor(c.terminal) });
        cy.edges().style({ 'line-color':c.edge, 'target-arrow-color':c.edge });
      });
      ['start','action','decision','invoke','terminal','edge'].forEach(k => {
        document.getElementById('sw-'+k).style.background = c[k];
        document.getElementById('c-'+k).value = c[k];
      });
    }

    let colors = { ...THEMES.default };
    applyColors(colors);

    ['start','action','decision','invoke','terminal','edge'].forEach(k => {
      document.getElementById('c-'+k).addEventListener('input', e => {
        colors[k] = e.target.value;
        applyColors(colors);
      });
    });

    document.querySelectorAll('.preset-btn').forEach(btn => {
      btn.addEventListener('click', () => {
        colors = { ...THEMES[btn.dataset.theme] };
        applyColors(colors);
      });
    });

    const colorPanel = document.getElementById('color-panel');
    const btnColors  = document.getElementById('btn-colors');
    const headerH    = document.querySelector('header').getBoundingClientRect().height;
    colorPanel.style.top  = headerH + 'px';
    detailPanel.style.top = headerH + 'px';

    function toggleColors() { colorPanel.classList.toggle('open'); btnColors.classList.toggle('active'); }
    btnColors.addEventListener('click', toggleColors);
    document.getElementById('btn-close-panel').addEventListener('click', () => {
      colorPanel.classList.remove('open');
      btnColors.classList.remove('active');
    });
  </script>
</body>
</html>"""


def to_html(elements: list[dict], title: str) -> str:
    safe_title       = title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    elements_json    = json.dumps(elements, ensure_ascii=False)
    col_headers_json = json.dumps(EXPECTED_COLUMNS, ensure_ascii=False)
    return (
        _HTML
        .replace("__TITLE__",           safe_title)
        .replace("__ELEMENTS_JSON__",   elements_json)
        .replace("__COL_HEADERS_JSON__", col_headers_json)
    )


# ── CLI ───────────────────────────────────────────────────────────────────────

def parse_args():
    p = argparse.ArgumentParser(
        description="Render a Cytoscape.js caller journey map from a CID_Search xlsx.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  python cid_journey.py CID-abc123_2026-03-04.xlsx\n"
            "  python cid_journey.py CID-abc123_2026-03-04.xlsx --output journey.html\n"
        ),
    )
    p.add_argument("xlsx", metavar="XLSX_FILE",
                   help="CID_Search output xlsx file")
    p.add_argument("--output", metavar="FILE",
                   help="Output HTML path (default: <input_stem>_journey.html)")
    return p.parse_args()


def main():
    args    = parse_args()
    in_path = Path(args.xlsx)

    if not in_path.exists():
        print(f"Error: file not found: {in_path}", file=sys.stderr)
        sys.exit(1)

    out_path = (
        Path(args.output)
        if args.output
        else in_path.with_name(in_path.stem + "_journey.html")
    )

    print(f"  Loading {in_path.name} …")
    rows = load_xlsx(in_path)
    if not rows:
        print("Error: no data rows found in xlsx.", file=sys.stderr)
        sys.exit(1)

    if len(rows) > 500:
        print(
            f"  Warning: {len(rows)} rows — journey map may render slowly in browser.",
            file=sys.stderr,
        )

    print(f"  Building journey map for {len(rows)} log event(s) …")
    elements = build_elements(rows)

    cid   = rows[0].get("ContactId", in_path.stem)
    title = f"Journey: {cid}"
    html  = to_html(elements, title)

    out_path.write_text(html, encoding="utf-8")
    print(f"  Journey map  \u2192  {out_path}")
    print("  Open in any browser (requires internet for Cytoscape CDN)\n")


if __name__ == "__main__":
    main()
